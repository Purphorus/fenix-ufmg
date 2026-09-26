"""Extrai texto dos materiais baixados (PDF, PPTX, DOCX, notebook, texto puro).

O texto vira um .txt ao lado do original, e o caminho vai para
`arquivos.texto_path`. É o insumo de `resumo.py` e `companion.py`.

A regra que orienta o módulo: **falhar visivelmente é melhor que produzir
texto vazio**. PDF escaneado (imagem, sem camada de texto) não gera .txt de
duas linhas — marca `precisa_ocr = 1` e aparece em `cli.py materiais` com o
selo correspondente. Resumo silenciosamente vazio é o pior resultado possível,
porque parece que funcionou.
"""

from __future__ import annotations

import json
import re
import sqlite3
import time
from dataclasses import dataclass
from pathlib import Path

import vetor

# Abaixo disto, um PDF é considerado escaneado em vez de "quase vazio".
# 40 caracteres por página passa folgado por slide de capa e por página de
# imagem legendada, e reprova página que só tem cabeçalho/rodapé do gerador.
MIN_CHARS_POR_PAGINA = 40

EXTENSOES = {".pdf", ".pptx", ".ppt", ".docx", ".doc", ".txt", ".md", ".ipynb"}


@dataclass
class Extracao:
    caminho: Path
    texto: str = ""
    paginas: int = 0
    precisa_ocr: bool = False
    erro: str = ""

    @property
    def ok(self) -> bool:
        return bool(self.texto) and not self.erro

    def __str__(self) -> str:
        if self.erro:
            return f"✗ {self.caminho.name}: {self.erro}"
        if self.precisa_ocr:
            return f"⚠ {self.caminho.name}: sem camada de texto — precisa de OCR"
        return f"✓ {self.caminho.name}: {len(self.texto)} caracteres, {self.paginas} pág."


# --------------------------------------------------------------------------
# Extratores por formato
# --------------------------------------------------------------------------


def _de_pdf(caminho: Path) -> Extracao:
    import pymupdf

    partes: list[str] = []
    with pymupdf.open(caminho) as doc:
        for n, pagina in enumerate(doc, 1):
            t = pagina.get_text("text").strip()
            if t:
                partes.append(f"\n--- página {n} ---\n{t}")
        paginas = doc.page_count

    texto = "\n".join(partes).strip()
    # Um PDF de slides escaneados costuma devolver só o número da página.
    escaneado = paginas > 0 and len(texto) < MIN_CHARS_POR_PAGINA * paginas
    return Extracao(caminho, "" if escaneado else texto, paginas, precisa_ocr=escaneado)


def _de_pptx(caminho: Path) -> Extracao:
    from pptx import Presentation

    pres = Presentation(caminho)
    partes = []
    for n, slide in enumerate(pres.slides, 1):
        linhas = [
            forma.text.strip()
            for forma in slide.shapes
            if getattr(forma, "has_text_frame", False) and forma.text.strip()
        ]
        # As notas do apresentador costumam ser o que o professor de fato disse.
        if slide.has_notes_slide:
            nota = (slide.notes_slide.notes_text_frame.text or "").strip()
            if nota:
                linhas.append(f"[notas do slide] {nota}")
        if linhas:
            partes.append(f"\n--- slide {n} ---\n" + "\n".join(linhas))
    return Extracao(caminho, "\n".join(partes).strip(), len(pres.slides))


def _de_docx(caminho: Path) -> Extracao:
    import docx

    doc = docx.Document(caminho)
    partes = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for tabela in doc.tables:  # cronograma de plano de ensino quase sempre é tabela
        for linha in tabela.rows:
            celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
            if celulas:
                partes.append(" | ".join(celulas))
    return Extracao(caminho, "\n".join(partes).strip(), 1)


def _de_notebook(caminho: Path) -> Extracao:
    nb = json.loads(caminho.read_text(encoding="utf-8", errors="replace"))
    partes = []
    for c in nb.get("cells", []):
        fonte = "".join(c.get("source", [])).strip()
        if not fonte:
            continue
        if c.get("cell_type") == "code":
            partes.append(f"```\n{fonte}\n```")
        else:
            partes.append(fonte)
    return Extracao(caminho, "\n\n".join(partes).strip(), len(nb.get("cells", [])))


def _de_texto(caminho: Path) -> Extracao:
    t = caminho.read_text(encoding="utf-8", errors="replace").strip()
    return Extracao(caminho, t, 1)


_EXTRATORES = {
    ".pdf": _de_pdf,
    ".pptx": _de_pptx,
    ".docx": _de_docx,
    ".ipynb": _de_notebook,
    ".txt": _de_texto,
    ".md": _de_texto,
}


def extrair(caminho: Path | str) -> Extracao:
    """Extrai o texto de um arquivo. Nunca levanta: erro vem no resultado.

    Extrair material é lote — um PPT corrompido no meio de 200 arquivos não
    pode derrubar a rodada inteira.
    """
    caminho = Path(caminho)
    if not caminho.is_file():
        return Extracao(caminho, erro="arquivo não encontrado")

    sufixo = caminho.suffix.lower()
    if sufixo in (".ppt", ".doc"):
        return Extracao(
            caminho,
            erro="formato antigo (.ppt/.doc); converta para .pptx/.docx",
        )
    extrator = _EXTRATORES.get(sufixo)
    if extrator is None:
        return Extracao(caminho, erro=f"sem extrator para {sufixo}")

    try:
        return extrator(caminho)
    except ImportError as e:
        return Extracao(caminho, erro=f"dependência faltando: {e}")
    except Exception as e:
        return Extracao(caminho, erro=f"{type(e).__name__}: {e}")


# --------------------------------------------------------------------------
# Integração com o banco
# --------------------------------------------------------------------------


def caminho_texto(original: Path) -> Path:
    """O .txt fica em textos/ ao lado do material, espelhando o nome."""
    return original.parent / "textos" / (original.stem + ".txt")


def extrair_arquivo(con: sqlite3.Connection, linha: sqlite3.Row) -> Extracao:
    """Extrai um registro de `arquivos` e atualiza texto_path / precisa_ocr."""
    res = extrair(linha["caminho_local"])

    destino = None
    if res.ok:
        destino = caminho_texto(Path(linha["caminho_local"]))
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_text(res.texto, encoding="utf-8")

    con.execute(
        "UPDATE arquivos SET texto_path = ?, precisa_ocr = ? WHERE id = ?",
        (str(destino) if destino else None, 1 if res.precisa_ocr else 0, linha["id"]),
    )
    indexar(con, linha["id"], linha["nome"], res.texto if res.ok else "")
    con.commit()
    return res


# --------------------------------------------------------------------------
# Fatiamento: do texto do arquivo para os trechos que serão citados
# --------------------------------------------------------------------------

# Os extratores de PDF e PPTX escrevem estes marcadores; é deles que sai o
# número de página/slide que aparece na citação.
_MARCADOR = re.compile(r"^--- (página|slide) (\d+) ---$", re.MULTILINE)

# Tamanho alvo do trecho. 1200 caracteres é ~300 tokens: cabe num resultado de
# busca sem estourar o orçamento e ainda carrega o raciocínio inteiro de um
# slide. Abaixo de 2000 uma página não é dividida — um slide é uma unidade de
# sentido, e cortá-lo ao meio separa o enunciado do gráfico que o explica.
ALVO_CHARS = 1200
MAX_CHARS = 2000
MIN_CHARS = 40          # abaixo disso é número de página solto, capa, rodapé

# Trecho de sumário/bibliografia casa com qualquer termo do curso e não
# responde nada. Não é excluído (às vezes é a única menção a um assunto), é
# rebaixado em `busca.py`.
_RUIDO = re.compile(
    r"^\s*(sum[áa]rio|[íi]ndice|refer[êe]ncias|bibliografia|agenda|roteiro)\b",
    re.IGNORECASE,
)


@dataclass
class Fatia:
    ordinal: int
    pagina: int | None
    rotulo: str          # "página 12" | "slide 4" | "" quando o formato não marca
    titulo: str
    texto: str

    @property
    def ruido(self) -> bool:
        return bool(_RUIDO.match(self.titulo or self.texto))


def _titulo(texto: str) -> str:
    """Primeira linha, se ela parecer título de slide e não frase cortada."""
    primeira = texto.strip().split("\n", 1)[0].strip()
    if 3 < len(primeira) <= 80 and not primeira.endswith((".", ",", ";", ":")):
        return primeira
    return ""


def _empacotar(texto: str) -> list[str]:
    """Quebra um bloco grande em pedaços de ~ALVO_CHARS, por parágrafo.

    Repete o último parágrafo no pedaço seguinte: sem sobreposição, uma
    definição que começa no fim de um pedaço e termina no começo do outro não
    é achada por nenhum dos dois.
    """
    if len(texto) <= MAX_CHARS:
        return [texto]

    paragrafos = [p for p in re.split(r"\n\s*\n", texto) if p.strip()]
    if len(paragrafos) < 2:  # parede de texto sem parágrafo: corta por linha
        paragrafos = [p for p in texto.split("\n") if p.strip()]

    pedacos: list[str] = []
    atual: list[str] = []
    n = 0
    for par in paragrafos:
        if atual and n + len(par) > ALVO_CHARS:
            pedacos.append("\n\n".join(atual))
            atual = [atual[-1]] if len(atual) > 1 else []
            n = sum(len(p) for p in atual)
        atual.append(par)
        n += len(par)
    if atual:
        pedacos.append("\n\n".join(atual))
    return pedacos


# Linha que aparece em pelo menos esta fração das páginas do arquivo é
# cabeçalho ou rodapé do gerador, não conteúdo. Medido: em Macro isso é 9% dos
# caracteres do curso e chega a 30% num arquivo; em Regional, 0%. Tirar antes
# de fatiar limpa a resposta E melhora o embedding, que hoje gasta dimensão
# representando o nome do departamento.
FRACAO_REPETIDA = 0.30
MIN_OCORRENCIAS = 3


def limpar_repetidas(texto: str) -> tuple[str, int]:
    """Remove cabeçalho e rodapé que se repetem pelo arquivo.

    Devolve (texto limpo, nº de linhas removidas). Conservador de propósito:
    exige a linha em 30% das páginas E pelo menos 3 vezes, então fórmula que
    reaparece em duas páginas sobrevive.
    """
    if not texto:
        return texto, 0
    paginas = _MARCADOR.split(texto)
    # split com 2 grupos de captura devolve [antes, rotulo, num, corpo, ...]
    corpos = paginas[3::3] if len(paginas) > 1 else [texto]
    if len(corpos) < MIN_OCORRENCIAS:
        return texto, 0

    from collections import Counter
    conta: Counter = Counter()
    for c in corpos:
        for l in {x.strip() for x in c.split("\n") if len(x.strip()) > 3}:
            conta[l] += 1

    limiar = max(MIN_OCORRENCIAS, FRACAO_REPETIDA * len(corpos))
    repetidas = {l for l, n in conta.items() if n >= limiar}
    if not repetidas:
        return texto, 0

    saida = []
    for linha in texto.split("\n"):
        if linha.strip() in repetidas:
            continue
        saida.append(linha)
    return "\n".join(saida), len(repetidas)


# Sinais de ruído medidos contra rótulo manual, com ECN300 como conjunto
# retido. O comprimento é o dominante: `n_chars < 200` pega 33% do ruído com
# 82% de precisão. Os demais somam pouco recall mas quase não erram.
RUIDO_CURTO = 200
_SUMARIO = re.compile(r"\.{5,}")
_SO_FONTE = re.compile(r"^\s*(fonte|source)\s*:", re.IGNORECASE)


def pontuar_ruido(f: "Fatia") -> float:
    """0 a 1: quanto o trecho parece boilerplate. Rebaixa, nunca exclui."""
    r = 0.0
    if len(f.texto) < RUIDO_CURTO:
        r += 0.5
    if f.ruido:                       # título de sumário/agenda/bibliografia
        r += 0.3
    if _SUMARIO.search(f.texto):      # pontilhado de índice
        r += 0.3
    if _SO_FONTE.match(f.texto):      # legenda de figura solta
        r += 0.3
    return min(r, 1.0)


def fatiar(texto: str) -> list[Fatia]:
    """Texto extraído -> trechos com âncora de página, prontos para indexar."""
    if not texto or not texto.strip():
        return []

    # (rotulo, numero, corpo) por marcador; sem marcador, o arquivo é um bloco.
    blocos: list[tuple[str, int | None, str]] = []
    marcas = list(_MARCADOR.finditer(texto))
    if not marcas:
        blocos.append(("", None, texto))
    else:
        if marcas[0].start() > MIN_CHARS:  # cabeçalho antes da primeira página
            blocos.append(("", None, texto[: marcas[0].start()]))
        for i, m in enumerate(marcas):
            fim = marcas[i + 1].start() if i + 1 < len(marcas) else len(texto)
            corpo = texto[m.end():fim]
            blocos.append((f"{m.group(1)} {m.group(2)}", int(m.group(2)), corpo))

    fatias: list[Fatia] = []
    for rotulo, pagina, corpo in blocos:
        for pedaco in _empacotar(corpo.strip()):
            pedaco = pedaco.strip()
            if len(pedaco) < MIN_CHARS:
                continue
            fatias.append(
                Fatia(len(fatias) + 1, pagina, rotulo, _titulo(pedaco), pedaco)
            )
    return fatias


# --------------------------------------------------------------------------
# Indexação: léxico (FTS5) e vetorial, no mesmo passo
# --------------------------------------------------------------------------


def indexar(con: sqlite3.Connection, arquivo_id: int, nome: str, texto: str) -> int:
    """(Re)indexa um arquivo: apaga os trechos antigos e grava os novos.

    Chamado a cada extração, então material republicado com conteúdo novo tem
    o índice trocado junto — sem passo separado que alguém possa esquecer.
    Texto vazio só remove o índice antigo.

    O vetor é chaveado pelo sha do trecho, então reindexar um arquivo que não
    mudou não paga o modelo de novo, e slide repetido entre duas aulas embute
    uma vez só. Sem `vetor.disponivel()`, grava só o lado léxico: a busca
    continua respondendo, com recall menor.
    """
    import hashlib

    antigos = [
        r[0] for r in con.execute(
            "SELECT id FROM trechos WHERE arquivo_id = ?", (arquivo_id,)
        ).fetchall()
    ]
    if antigos:
        marcas = ",".join("?" * len(antigos))
        con.execute(f"DELETE FROM trechos_fts WHERE trecho_id IN ({marcas})", antigos)
        con.execute("DELETE FROM trechos WHERE arquivo_id = ?", (arquivo_id,))

    # Cabeçalho e rodapé saem ANTES de fatiar: eles inflam o trecho, sujam a
    # citação e ainda entram no embedding como se fossem assunto.
    texto, _removidas = limpar_repetidas(texto)
    fatias = fatiar(texto)
    if not fatias:
        return 0

    shas = []
    itens = []
    secao = _secao_do_arquivo(con, arquivo_id)
    titulo_vigente = ""
    for f in fatias:
        sha = hashlib.sha256(f.texto.encode("utf-8")).hexdigest()
        shas.append(sha)
        titulo_vigente = f.titulo or titulo_vigente
        ctx = contexto_do_trecho(nome, secao, titulo_vigente, f.texto)
        itens.append((sha, f.texto, ctx))
        cur = con.execute(
            """INSERT INTO trechos
               (arquivo_id, ordinal, pagina, rotulo, titulo, texto, n_chars,
                sha256, ruido)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            (arquivo_id, f.ordinal, f.pagina, f.rotulo, f.titulo,
             f.texto, len(f.texto), sha, pontuar_ruido(f)),
        )
        con.execute(
            "INSERT INTO trechos_fts (nome, texto, trecho_id) VALUES (?,?,?)",
            (_nome_fts(con, nome, ctx), f.texto, cur.lastrowid),
        )

    embutir_faltantes(con, itens)
    return len(fatias)


# --------------------------------------------------------------------------
# Contexto do trecho (contextual retrieval, versão sem LLM)
# --------------------------------------------------------------------------
#
# Um trecho do meio de um deck não diz de que aula é: "a curva desloca para a
# direita" pode ser oferta agregada ou demanda por moeda. A Anthropic mediu
# que pôr um contexto curto na frente do trecho, antes de embutir e de indexar
# no BM25, corta a falha de recuperação em 35% (só embedding) e 49% (os dois).
# Lá o contexto é escrito por um LLM; aqui ele sai do que já se sabe sem gastar
# token: nome do arquivo, seção do Moodle e o último título de slide/página
# visto até o trecho. No perfilador a mesma ideia rendeu 9 → 12 de 20.
#
# O texto do trecho e o sha NÃO mudam — rótulos, cobertura e citação são
# chaveados por eles (invariante 6). O contexto só entra no que a busca
# compara, e liga por banco em `config_busca.contexto`.

_EXTENSAO = re.compile(r"\.(pdf|pptx?|docx?|txt|md|ipynb)$", re.IGNORECASE)


def contexto_do_trecho(nome: str, secao: str, titulo: str, texto: str) -> str:
    """Uma ou duas linhas que situam o trecho. Vazio quando não há o que dizer."""
    partes = []
    n = _EXTENSAO.sub("", (nome or "").strip())
    if n:
        partes.append(n)
    if secao and secao.strip() and secao.strip() not in n:
        partes.append(secao.strip())
    linha = " — ".join(partes)
    # Título que o próprio trecho já abre não se repete: repetir dobraria o
    # peso dele no embedding sem acrescentar nada.
    t = (titulo or "").strip()
    if t and not (texto or "").lstrip().startswith(t):
        linha = f"{linha}\n{t}" if linha else t
    return linha


def _secao_do_arquivo(con: sqlite3.Connection, arquivo_id: int) -> str:
    r = con.execute("SELECT secao FROM arquivos WHERE id = ?", (arquivo_id,)).fetchone()
    return (r[0] if r else "") or ""


def _nome_fts(con: sqlite3.Connection, nome: str, ctx: str) -> str:
    """O que vai na coluna `nome` do FTS: o nome, ou o contexto inteiro.

    O nome do arquivo já era indexado; com o contexto ligado, seção e título
    vigente entram na mesma coluna — é a metade BM25 do contextual retrieval.
    """
    if vetor.contexto_ligado(con) and ctx:
        return ctx
    return nome or ""


def _com_contexto(ctx: str, texto: str) -> str:
    return f"{ctx}\n\n{texto}" if ctx else texto


def _chaves(con: sqlite3.Connection) -> list[tuple[str, bool]]:
    """Os vetores que cada trecho precisa: (chave, com_contexto).

    Sempre o do texto puro, que o grafo usa para achar conteúdo repetido entre
    arquivos; e, com o contexto ligado, o que a busca compara com a consulta.
    """
    base = vetor.chave_conteudo(con)
    busca_ = vetor.chave(con)
    return [(base, False)] + ([(busca_, True)] if busca_ != base else [])


def embutir_faltantes(con: sqlite3.Connection, itens: list[tuple]) -> int:
    """Embute os trechos que ainda não têm vetor para o modelo atual.

    `itens` são (sha, texto) ou (sha, texto, contexto).
    """
    if not itens or not vetor.disponivel(con):
        return 0
    m = vetor.modelo(con)
    agora = time.strftime("%Y-%m-%dT%H:%M:%S")
    total = 0
    for chave, com_ctx in _chaves(con):
        pendentes_: dict[str, str] = {}
        for it in itens:
            sha, texto = it[0], it[1]
            ctx = it[2] if len(it) > 2 else ""
            if sha in pendentes_:
                continue
            ja = con.execute(
                "SELECT 1 FROM vetores WHERE sha256 = ? AND modelo = ?", (sha, chave)
            ).fetchone()
            if not ja:
                pendentes_[sha] = _com_contexto(ctx, texto) if com_ctx else texto
        if not pendentes_:
            continue
        shas = list(pendentes_)
        blobs = vetor.embutir([pendentes_[s] for s in shas], m)
        if not blobs:
            continue
        con.executemany(
            """INSERT OR REPLACE INTO vetores (sha256, modelo, dim, vetor, criado_em)
               VALUES (?,?,?,?,?)""",
            [(s, chave, m.dim, b, agora) for s, b in zip(shas, blobs)],
        )
        total += len(blobs)
    return total


def vetores_pendentes(con: sqlite3.Connection) -> int:
    """Quantos trechos indexados ainda não têm algum vetor do modelo atual."""
    n = 0
    for chave, _ in _chaves(con):
        n = max(n, con.execute(
            """SELECT COUNT(*) FROM (
                   SELECT DISTINCT t.sha256 FROM trechos t
                   LEFT JOIN vetores v ON v.sha256 = t.sha256 AND v.modelo = ?
                   WHERE v.sha256 IS NULL)""",
            (chave,),
        ).fetchone()[0])
    return n


def _itens_com_contexto(con: sqlite3.Connection, shas: set[str] | None = None) -> list[tuple]:
    """(sha, texto, contexto) de cada trecho, com o título vigente por arquivo."""
    itens, atual, titulo = [], None, ""
    for r in con.execute(
        """SELECT t.sha256, t.texto, t.titulo, t.arquivo_id, a.nome, a.secao
           FROM trechos t JOIN arquivos a ON a.id = t.arquivo_id
           ORDER BY t.arquivo_id, t.ordinal"""
    ):
        if r["arquivo_id"] != atual:
            atual, titulo = r["arquivo_id"], ""
        titulo = r["titulo"] or titulo
        if shas is None or r["sha256"] in shas:
            itens.append((r["sha256"], r["texto"],
                          contexto_do_trecho(r["nome"], r["secao"] or "", titulo, r["texto"])))
    return itens


BLOCO_VETORES = 64


def completar_vetores(con: sqlite3.Connection, limite: int = 0) -> int:
    """Embute os trechos que ficaram sem vetor. Roda junto da extração.

    Existe porque a extração pode acontecer com o modelo indisponível — o
    agente do launchd rodando sem rede, ou antes do primeiro download. Sem
    isto, aquele arquivo ficaria fora do lado semântico para sempre, e em
    silêncio: a busca continuaria respondendo, só que sem ele. Aqui o índice
    se conserta na rodada seguinte, sozinho — inclusive depois de trocar o
    modelo ou ligar o contexto em `config_busca`.
    """
    if not vetor.disponivel(con):
        return 0
    faltam: set[str] = set()
    for chave, _ in _chaves(con):
        faltam |= {r[0] for r in con.execute(
            """SELECT DISTINCT t.sha256 FROM trechos t
               LEFT JOIN vetores v ON v.sha256 = t.sha256 AND v.modelo = ?
               WHERE v.sha256 IS NULL""", (chave,))}
    if not faltam:
        return 0
    itens = _itens_com_contexto(con, faltam)
    if limite:
        itens = itens[:limite]
    # Em blocos, gravando a cada um: trocar para um modelo grande embute o
    # corpus inteiro, leva minutos, e uma interrupção no meio não pode jogar
    # fora o que já foi feito — a rodada seguinte continua de onde parou.
    n = 0
    for i in range(0, len(itens), BLOCO_VETORES):
        feitos = embutir_faltantes(con, itens[i:i + BLOCO_VETORES])
        if feitos:
            con.commit()
            n += feitos
    return n


def reindexar_tudo(con: sqlite3.Connection) -> tuple[int, int]:
    """Reconstrói o índice a partir dos .txt já extraídos.

    Devolve (arquivos, trechos). É o que roda depois de mudar o fatiamento ou
    de trocar o modelo de embedding.
    """
    con.execute("DELETE FROM trechos_fts")
    con.execute("DELETE FROM trechos")
    con.execute("DELETE FROM grafo_documentos")
    arquivos = trechos = 0
    for r in con.execute(
        "SELECT id, nome FROM arquivos WHERE texto_path IS NOT NULL"
    ).fetchall():
        t = texto_de(con, r["id"])
        if t:
            n = indexar(con, r["id"], r["nome"], t)
            if n:
                arquivos += 1
                trechos += n
    # Vetor de sha que não é mais de nenhum trecho vira lixo permanente: a
    # limpeza do texto muda o sha de todo mundo, e sem isto o banco cresce a
    # cada reindexação.
    con.execute(
        """DELETE FROM vetores WHERE NOT EXISTS
           (SELECT 1 FROM trechos t WHERE t.sha256 = vetores.sha256)"""
    )
    con.commit()
    _reconstruir_grafo(con)
    return arquivos, trechos


def _reconstruir_grafo(con: sqlite3.Connection) -> int:
    """Refaz o grafo entre documentos. Erro aqui não derruba a indexação.

    O grafo é derivado dos vetores, então ele nasce e morre com eles: sem
    modelo de embedding não há grafo, e a busca continua respondendo — só
    perde a travessia entre documentos.
    """
    try:
        import grafo
        return grafo.construir_tudo(con)
    except Exception as e:  # noqa: BLE001 — derivado, nunca essencial
        print(f"  ! grafo de documentos: {type(e).__name__}: {e}")
        return 0


def pendentes(con: sqlite3.Connection, site: str | None = None, refazer: bool = False):
    """Arquivos baixados que ainda não têm texto extraído.

    `sync.py` zera `texto_path` quando o conteúdo de um arquivo muda, então
    material republicado com alteração real volta a aparecer aqui sozinho.
    """
    sql = "SELECT * FROM arquivos WHERE caminho_local IS NOT NULL"
    params: list = []
    if not refazer:
        # precisa_ocr = 0 evita reprocessar escaneado a cada rodada, sem OCR
        # não vai mudar de resultado.
        sql += " AND texto_path IS NULL AND precisa_ocr = 0"
    if site:
        sql += " AND site = ?"
        params.append(site)
    return con.execute(sql + " ORDER BY id", params).fetchall()


@dataclass
class ResultadoExtracao:
    extraidos: int = 0
    ocr: int = 0
    erros: int = 0
    vetores: int = 0
    arestas: int = 0
    detalhes: list[str] = None  # type: ignore[assignment]

    def __post_init__(self):
        if self.detalhes is None:
            self.detalhes = []


def extrair_pendentes(
    con: sqlite3.Connection,
    site: str | None = None,
    refazer: bool = False,
    limite: int | None = None,
) -> ResultadoExtracao:
    res = ResultadoExtracao()
    linhas = pendentes(con, site, refazer)
    if limite:
        linhas = linhas[:limite]

    for linha in linhas:
        e = extrair_arquivo(con, linha)
        if e.erro:
            res.erros += 1
        elif e.precisa_ocr:
            res.ocr += 1
        else:
            res.extraidos += 1
        res.detalhes.append(str(e))

    # Rede caída ou modelo ainda não baixado deixam trecho sem vetor; a rodada
    # seguinte completa. Erro aqui não derruba a extração, que já terminou.
    try:
        res.vetores = completar_vetores(con)
    except Exception as e:  # noqa: BLE001 — lote não pode cair por causa disto
        res.detalhes.append(f"✗ vetores: {type(e).__name__}: {e}")
    # Material novo muda a vizinhança entre documentos; sem isto o grafo
    # envelhece em silêncio e a busca passa a atravessar para o lugar errado.
    if res.extraidos or res.vetores:
        res.arestas = _reconstruir_grafo(con)
    return res


def texto_de(con: sqlite3.Connection, arquivo_id: int) -> str:
    """Lê o texto já extraído de um arquivo (vazio se não houver)."""
    r = con.execute(
        "SELECT texto_path FROM arquivos WHERE id = ?", (arquivo_id,)
    ).fetchone()
    if not r or not r["texto_path"]:
        return ""
    p = Path(r["texto_path"])
    return p.read_text(encoding="utf-8") if p.is_file() else ""
